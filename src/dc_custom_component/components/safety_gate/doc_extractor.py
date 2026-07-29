"""Safety Gate document extraction — deepset port of `preprocessing/doc_extract.py`.

One PDF or XLSX in, one Haystack Document out whose `content` is byte-comparable
with the reference `<alertId>/output/<doc>.md`.

Text layer first, pixels only as a fallback
-------------------------------------------
PDFs are read via pypdfium2's embedded text layer. A page is rasterised and OCR'd
only when it fails one of two gates:

  * **quantity** — fewer than 50 usable (alphanumeric) characters: a scan
  * **quality**  — enough characters but under 6 real words per 100, the
    signature of a broken CID font that emits stray alphanumerics with no words

So a clean 100-page digital report costs nothing, while a 4-page scan costs four
OCR passes. The frontmatter records which path was taken (`embedded` / `ocr` /
`mixed`), and the quality gate adds a `note:` naming the pages it rescued.

XLSX never touches OCR — openpyxl to markdown tables, capped at 500 rows x 40
columns.

Port decisions
--------------
**`artifact: doc` goes in `meta` but NOT in the markdown.** The local renderer
emits no `artifact` field for documents (unlike image and barcode markdown), so
adding one to `content` would break byte-parity with the reference. deepset has no
filenames to distinguish document kinds by, so the field is required for
retrieval — it therefore lives in `meta` only. This resolves a genuine conflict
between the two acceptance criteria in issue #37.

**`generated:` is unquoted here.** `doc_extract.write_markdown` writes it bare
while `ocr_render_md` quotes it. Preserved deliberately; it is a real difference
between the two reference formats, not an inconsistency to tidy up.

**`## Page N` headings are load-bearing.** `optimization/data.py::_PAGE_SPLIT`
splits on that exact form to do verdict-aware truncation. Changing the heading
silently disables page-aware truncation for 100-page lab reports.

**The OCR reader is built lazily and is primary-only.** Mirrors the local
`readers_box` pattern — most alerts have clean text layers and never construct a
recognizer. Note the local code uses `readers.primary` for scanned pages with **no
Cyrillic escalation**, unlike the image path; that asymmetry is preserved rather
than silently improved.
"""

from __future__ import annotations

import os

# MUST precede any paddle/numpy import — see paddle_ocr.py for the reasoning.
for _var in (
    "OMP_NUM_THREADS",
    "MKL_NUM_THREADS",
    "OPENBLAS_NUM_THREADS",
    "NUMEXPR_NUM_THREADS",
    "VECLIB_MAXIMUM_THREADS",
    "OMP_THREAD_LIMIT",
):
    os.environ.setdefault(_var, "1")
os.environ.setdefault("GLOG_minloglevel", "3")
os.environ.setdefault("FLAGS_call_stack_level", "0")
os.environ.setdefault("PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK", "True")

import hashlib  # noqa: E402
import logging  # noqa: E402
import re  # noqa: E402
import tempfile  # noqa: E402
from datetime import datetime  # noqa: E402
from pathlib import Path  # noqa: E402
from typing import Any, Dict, List, Optional, Tuple, Union, cast  # noqa: E402

from haystack import (  # noqa: E402
    Document,
    component,
    default_from_dict,
    default_to_dict,
)
from haystack.components.converters.utils import (  # noqa: E402
    get_bytestream_from_source,
    normalize_metadata,
)
from haystack.dataclasses import ByteStream  # noqa: E402

# Identical quoting in both reference renderers — imported rather than copied so
# the two component families cannot drift apart.
from dc_custom_component.components.safety_gate.paddle_ocr import (  # noqa: E402
    DET_MODEL,
    MAX_SIDE,
    PRIMARY_REC,
    parse_prefixed_filename,
    yaml_str,
)

logger = logging.getLogger(__name__)

# --- constants, all mirrored from preprocessing/doc_extract.py ---------------
DOC_EXTS = {".pdf", ".xlsx"}
# Below this many usable characters, a PDF page is treated as scanned. Length
# alone is not enough: some text layers are entirely Private-Use-Area codepoints
# (invisible) or raw glyph indexes.
MIN_TEXT_CHARS = 50
# An OCR'd fallback page under this many alnum chars gets one higher-resolution
# retry before being accepted as (near-)blank.
MIN_OCR_CHARS = 20
# PDF points are 72 dpi, so 150/72 renders at ~150 dpi.
RENDER_SCALE = 150 / 72
RETRY_SCALE = 300 / 72
MAX_SHEET_ROWS = 500
MAX_SHEET_COLS = 40
# Real-word density below which a page's text layer is treated as garbled. Clean
# prose in any EU language scores ~14-20 words per 100 usable chars; a broken CID
# font scores ~0-2 while still carrying plenty of stray alphanumerics.
GARBLED_MAX_DENSITY = 6.0

_WORD_RX = re.compile(r"[A-Za-zÀ-žΑ-ωА-я]{3,}")


# --- pure helpers, ported 1:1 from doc_extract.py -----------------------------


def sanitize_name(stem: str) -> str:
    """Note the "document" fallback — `ocr_render_md` uses "image" instead."""
    keep = [
        ch if (ch.isalnum() or ch in "-_ .()+&") else "_" for ch in str(stem).strip()
    ]
    return " ".join("".join(keep).split()) or "document"


def usable_count(text: str) -> int:
    """Alphanumeric characters in any script — the measure of real content."""
    return sum(1 for ch in text if ch.isalnum())


def word_count(text: str) -> int:
    return len(_WORD_RX.findall(text))


def looks_garbled(
    text: str,
    min_chars: int = MIN_TEXT_CHARS,
    max_density: float = GARBLED_MAX_DENSITY,
) -> bool:
    """Enough characters to look like content, almost none forming real words.

    Thresholds are parameters so the component's configurable values are used on
    every call — including the post-OCR retry check, which would otherwise
    silently fall back to the module defaults.
    """
    usable = usable_count(text)
    if usable < min_chars:
        return False  # too short to judge; the quantity gate rules
    return word_count(text) * 100 / usable < max_density


def clean_page_text(raw: str) -> str:
    """Repair and scrub one page's extracted text layer.

    Some PDFs emit their whole text as Private-Use-Area codepoints
    (0xF000 + ascii, from symbolic font cmaps). Decode that when it dominates the
    page AND the decoding actually improves things. Then drop C0 control
    characters (glyph indexes from fonts without ToUnicode) and leftover PUA
    symbols, which have no visible glyph and would render as scrambled output.
    """
    non_ws = [c for c in raw if not c.isspace()]
    if non_ws:
        pua = sum(1 for c in non_ws if 0xF000 <= ord(c) <= 0xF0FF)
        if pua / len(non_ws) > 0.5:
            decoded = "".join(
                chr(ord(c) - 0xF000) if 0xF000 <= ord(c) <= 0xF0FF else c for c in raw
            )
            if usable_count(decoded) > usable_count(raw):
                raw = decoded
    return "".join(
        c
        for c in raw
        if not (ord(c) < 32 and c not in "\t\n\r") and not (0xE000 <= ord(c) <= 0xF8FF)
    )


def cell_str(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace("\n", " ").replace("|", "\\|").strip()


def render_document_md(
    title: str,
    alert_id: str,
    folder: str,
    file_name: str,
    unit_name: str,
    n_units: int,
    kind: str,
    note: Optional[str],
    sections: List[Tuple[str, str]],
    engine: str,
    generated: Optional[str] = None,
) -> str:
    """Byte-for-byte `doc_extract.write_markdown`.

    `generated` is deliberately UNQUOTED, and there is deliberately no
    `artifact:` line — both match the reference renderer exactly.
    """
    frontmatter = [
        "---",
        f"title: {yaml_str(title)}",
        f"alert_id: {alert_id}",
        f"source: {yaml_str(f'{folder}/{file_name}')}",
        f"folder: {folder}",
        f"file_type: {Path(file_name).suffix.lower().lstrip('.')}",
        f"{unit_name}: {n_units}",
        f"extraction: {kind}",
        f"generated: {generated or datetime.now().isoformat(timespec='seconds')}",
        f"engine: {yaml_str(engine)}",
    ]
    if note:
        frontmatter.append(f"note: {yaml_str(note)}")
    frontmatter.append("---")

    body = [f"\n# {title}"]
    for heading, text in sections:
        body.append(f"\n## {heading}\n")
        body.append(text if text else "_(no text)_")
    return "\n".join(frontmatter) + "\n" + "\n".join(body) + "\n"


def document_id(alert_id: str, folder: str, file_name: str) -> str:
    """Stable id so re-indexing overwrites rather than accumulating. Excludes
    content, which carries a wall-clock `generated` timestamp."""
    return hashlib.sha1(
        f"safety-gate-doc|{alert_id}|{folder}|{file_name}".encode("utf-8")
    ).hexdigest()


@component
class SafetyGateDocExtractor:
    """Extracts text from Safety Gate alert PDFs and spreadsheets into markdown.

    :param enable_mkldnn: leave False. True crashes paddle's text detection on
        deepset's Linux/x86_64 (measured 2026-07-28) — same constraint as the
        image OCR component.
    :param min_text_chars: usable-character threshold below which a PDF page is
        treated as scanned and sent to OCR.
    :param garbled_max_density: real words per 100 usable characters below which
        a text layer is treated as broken and sent to OCR anyway.
    :param default_folder / default_alert_id: used when a source carries neither
        metadata nor a `<alertId>__<folder>__` filename prefix.
    """

    def __init__(
        self,
        enable_mkldnn: bool = False,
        min_text_chars: int = MIN_TEXT_CHARS,
        garbled_max_density: float = GARBLED_MAX_DENSITY,
        max_sheet_rows: int = MAX_SHEET_ROWS,
        max_sheet_cols: int = MAX_SHEET_COLS,
        max_side: int = MAX_SIDE,
        default_folder: str = "published",
        default_alert_id: str = "",
    ) -> None:
        self.enable_mkldnn = enable_mkldnn
        self.min_text_chars = min_text_chars
        self.garbled_max_density = garbled_max_density
        self.max_sheet_rows = max_sheet_rows
        self.max_sheet_cols = max_sheet_cols
        self.max_side = max_side
        self.default_folder = default_folder
        self.default_alert_id = default_alert_id

        self._reader: Any = None
        self._engines: Dict[str, str] = {}

    def to_dict(self) -> Dict[str, Any]:
        return cast(
            Dict[str, Any],
            default_to_dict(
                self,
                enable_mkldnn=self.enable_mkldnn,
                min_text_chars=self.min_text_chars,
                garbled_max_density=self.garbled_max_density,
                max_sheet_rows=self.max_sheet_rows,
                max_sheet_cols=self.max_sheet_cols,
                max_side=self.max_side,
                default_folder=self.default_folder,
                default_alert_id=self.default_alert_id,
            ),
        )

    @classmethod
    def from_dict(cls, data: Dict[str, Any]) -> "SafetyGateDocExtractor":
        return cast("SafetyGateDocExtractor", default_from_dict(cls, data))

    # --- lifecycle ----------------------------------------------------------

    def warm_up(self) -> None:
        """Resolve the engine strings only.

        No OCR model is built here: the engine string appears in EVERY document's
        frontmatter, including pure-embedded ones, but the recognizer is needed
        only for scanned pages — so it stays lazy, exactly as the local
        `readers_box` does.
        """
        if self._engines:
            return
        from importlib.metadata import version

        import paddleocr

        self._engines = {
            "pdf": (
                f"pypdfium2 {version('pypdfium2')} + paddleocr "
                f"{paddleocr.__version__} (scanned pages) +qgate"
            ),
            "xlsx": f"openpyxl {version('openpyxl')}",
        }

    @property
    def _ocr(self) -> Any:
        """Primary recognizer, built on first scanned page. No Cyrillic
        escalation — the local document path does not have it either."""
        if self._reader is None:
            from paddleocr import PaddleOCR

            self._reader = PaddleOCR(
                text_detection_model_name=DET_MODEL,
                text_recognition_model_name=PRIMARY_REC,
                use_doc_orientation_classify=False,
                use_doc_unwarping=False,
                use_textline_orientation=True,
                device="cpu",
                enable_mkldnn=self.enable_mkldnn,
                cpu_threads=1,
            )
        return self._reader

    # --- PDF ----------------------------------------------------------------

    def _looks_garbled(self, text: str) -> bool:
        """`looks_garbled` bound to this instance's configured thresholds."""
        return looks_garbled(text, self.min_text_chars, self.garbled_max_density)

    def _ocr_page_bitmap(self, page: Any, scale: float = RENDER_SCALE) -> str:
        """Render one PDF page and OCR it. Port of `ocr_page_bitmap`."""
        import numpy as np

        array = page.render(scale=scale).to_numpy()  # HxWx{3,4} RGB(A)
        if array.ndim == 3 and array.shape[2] == 4:
            array = array[:, :, :3]
        height, width = array.shape[:2]
        if max(height, width) > self.max_side:
            from PIL import Image

            factor = self.max_side / max(height, width)
            resized = Image.fromarray(array).resize(
                (max(1, int(width * factor)), max(1, int(height * factor))),
                Image.LANCZOS,
            )
            array = np.asarray(resized)
        result = self._ocr.predict(array[:, :, ::-1].copy())[0]  # RGB -> BGR
        return "\n".join(list(result.get("rec_texts") or []))

    def _extract_pdf(
        self, path: Path
    ) -> Tuple[List[Tuple[str, str]], int, str, Optional[str]]:
        """Port of `extract_pdf`. -> (sections, n_pages, kind, note)."""
        import pypdfium2 as pdfium

        document = pdfium.PdfDocument(str(path))
        sections: List[Tuple[str, str]] = []
        methods: set = set()
        qgate_pages: List[str] = []
        try:
            n_pages = len(document)
            for index in range(n_pages):
                page = document[index]
                raw = page.get_textpage().get_text_bounded() or ""
                text = clean_page_text(raw).strip()

                # QUALITY gate on top of the QUANTITY gate: a broken text layer
                # can carry plenty of stray alphanumerics yet zero real words,
                # and must fall through to OCR despite passing the char count.
                gate_garbled = self._looks_garbled(text)
                if usable_count(text) >= self.min_text_chars and not gate_garbled:
                    methods.add("embedded")
                else:
                    ocr_text = self._ocr_page_bitmap(page)
                    if usable_count(ocr_text) < MIN_OCR_CHARS or (
                        gate_garbled and self._looks_garbled(ocr_text)
                    ):
                        retry = self._ocr_page_bitmap(page, scale=RETRY_SCALE)
                        if usable_count(retry) > usable_count(ocr_text):
                            ocr_text = retry
                    # On garble-gated pages keep the better reading by word
                    # count — never worse than the embedded mojibake. Pages that
                    # arrived via the quantity gate keep the original behaviour.
                    if ocr_text.strip() and (
                        not gate_garbled or word_count(ocr_text) >= word_count(text)
                    ):
                        text = ocr_text.strip()
                        methods.add("ocr")
                        if gate_garbled:
                            qgate_pages.append(str(index + 1))
                    elif text:
                        methods.add("embedded")
                sections.append((f"Page {index + 1}", text))
        finally:
            document.close()

        if not methods:
            kind = "none"
        elif methods == {"embedded"}:
            kind = "embedded"
        elif methods == {"ocr"}:
            kind = "ocr"
        else:
            kind = "mixed"
        note = (
            f"quality-gate: garbled embedded text on page(s) "
            f"{', '.join(qgate_pages)} -> OCR fallback"
            if qgate_pages
            else None
        )
        return sections, n_pages, kind, note

    # --- XLSX ---------------------------------------------------------------

    def _extract_xlsx(
        self, path: Path
    ) -> Tuple[List[Tuple[str, str]], int, str, Optional[str]]:
        """Port of `extract_xlsx`. Each sheet becomes a markdown table."""
        import openpyxl

        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        sections: List[Tuple[str, str]] = []
        notes: List[str] = []
        for sheet in workbook.worksheets:
            rows: List[List[str]] = []
            truncated = False
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_index >= self.max_sheet_rows:
                    truncated = True
                    break
                cells = [cell_str(c) for c in row[: self.max_sheet_cols]]
                if any(cells):
                    rows.append(cells)
            if not rows:
                sections.append((f"Sheet: {sheet.title}", "_(empty sheet)_"))
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            lines = [
                "| " + " | ".join(rows[0]) + " |",
                "| " + " | ".join(["---"] * width) + " |",
            ]
            for row_cells in rows[1:]:
                lines.append("| " + " | ".join(row_cells) + " |")
            if truncated:
                lines.append(f"\n_(truncated at {self.max_sheet_rows} rows)_")
                notes.append(f"sheet '{sheet.title}' truncated")
            sections.append((f"Sheet: {sheet.title}", "\n".join(lines)))
        workbook.close()
        return sections, len(sections), "spreadsheet", "; ".join(notes) or None

    # --- provenance ---------------------------------------------------------

    def _resolve(self, stream: ByteStream, extra: Dict[str, Any]) -> Dict[str, str]:
        """Same four-source resolution as the OCR component: explicit metadata,
        then a `<alertId>__<folder>__` filename prefix, then a real directory
        path, then defaults. The prefix is always stripped from the name."""
        meta: Dict[str, Any] = {**(stream.meta or {}), **extra}

        file_path = str(meta.get("file_path") or "")
        name = str(meta.get("file_name") or (Path(file_path).name if file_path else ""))
        if not name:
            name = "document"

        prefix_alert, prefix_folder, name = parse_prefixed_filename(name)

        folder = str(meta.get("folder") or "") or prefix_folder
        if not folder and file_path:
            parts = Path(file_path).parts
            for candidate in ("published", "restricted"):
                if candidate in parts:
                    folder = candidate
                    break
        folder = folder or self.default_folder

        alert_id = str(meta.get("alert_id") or "") or prefix_alert
        if not alert_id and file_path:
            parts = Path(file_path).parts
            for index, part in enumerate(parts):
                if part in ("published", "restricted") and index > 0:
                    alert_id = parts[index - 1]
                    break
        alert_id = alert_id or self.default_alert_id or "unknown"

        return {"alert_id": alert_id, "folder": folder, "file_name": name}

    # --- run ----------------------------------------------------------------

    @component.output_types(documents=List[Document])
    def run(
        self,
        sources: List[Union[str, Path, ByteStream]],
        meta: Optional[Union[Dict[str, Any], List[Dict[str, Any]]]] = None,
    ) -> Dict[str, List[Document]]:
        """Extract each source document into one markdown Document.

        An extraction failure produces a Document with `extraction: error` and
        the exception in `note:`, exactly as the local script does — the failure
        stays visible downstream instead of the document silently not existing.
        """
        self.warm_up()

        documents: List[Document] = []
        metadata = normalize_metadata(meta, sources_count=len(sources))

        for source, extra in zip(sources, metadata):
            try:
                stream = get_bytestream_from_source(source)
            except Exception as exc:
                logger.warning("Could not read source %s: %s", source, exc)
                continue

            resolved = self._resolve(stream, extra)
            name = resolved["file_name"]
            suffix = Path(name).suffix.lower()
            if suffix not in DOC_EXTS:
                logger.warning("Skipping %s: not a PDF or XLSX", name)
                continue

            documents.append(self._process(stream, resolved, extra, suffix))

        return {"documents": documents}

    def _process(
        self,
        stream: ByteStream,
        resolved: Dict[str, str],
        extra: Dict[str, Any],
        suffix: str,
    ) -> Document:
        name = resolved["file_name"]
        title = sanitize_name(Path(name).stem)

        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as handle:
            handle.write(stream.data)
            staged = Path(handle.name)
        try:
            if suffix == ".pdf":
                sections, n_units, kind, note = self._extract_pdf(staged)
                unit_name, engine = "pages", self._engines["pdf"]
            else:
                sections, n_units, kind, note = self._extract_xlsx(staged)
                unit_name, engine = "sheets", self._engines["xlsx"]
        except Exception as exc:
            # Mirrors doc_extract.process_alert's except branch: an error
            # document, not a missing one.
            sections, n_units, kind = [], 0, "error"
            note = f"{type(exc).__name__}: {exc}"
            unit_name, engine = "pages", self._engines.get("pdf", "unknown")
            logger.warning("Extraction failed for %s: %s", name, exc)
        finally:
            staged.unlink(missing_ok=True)

        content = render_document_md(
            title=title,
            alert_id=resolved["alert_id"],
            folder=resolved["folder"],
            file_name=name,
            unit_name=unit_name,
            n_units=n_units,
            kind=kind,
            note=note,
            sections=sections,
            engine=engine,
        )

        document_meta: Dict[str, Any] = {
            **{k: v for k, v in (stream.meta or {}).items() if k != "file_path"},
            **extra,
            "alert_id": resolved["alert_id"],
            # NOT written into the markdown — see the module docstring. deepset
            # has no filenames, so retrieval needs this field in meta.
            "artifact": "doc",
            "source": f"{resolved['folder']}/{name}",
            "folder": resolved["folder"],
            "file_name": name,
            "file_type": suffix.lstrip("."),
            unit_name: n_units,
            "extraction": kind,
            "engine": engine,
        }
        if note:
            document_meta["note"] = note

        return Document(
            id=document_id(resolved["alert_id"], resolved["folder"], name),
            content=content,
            meta=document_meta,
        )
